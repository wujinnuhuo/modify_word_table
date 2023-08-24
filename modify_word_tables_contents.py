import os
from docx import Document


def replace_specific_cell_content(table, row_index, col_index, new_content):
    if row_index < len(table.rows) and col_index < len(table.columns):
        cell = table.cell(row_index, col_index)
        cell.text = new_content


def process_files(file_time_mapping, start_date_row, start_date_col, start_date_content,
                  trigger_time_row, trigger_time_col,
                  venue_name_row, venue_name_col, venue_name_content,
                  project_name_row, project_name_col, project_name_content, output_folder):
    for file_name, trigger_time_content in file_time_mapping.items():
        # 搜索并读取Word文件
        for root, dirs, files in os.walk(".", topdown=True):
            if file_name in files:
                docx_path = os.path.join(root, file_name)
                doc = Document(docx_path)

                # 在每个表格中查找并替换字段
                for table in doc.tables:
                    replace_specific_cell_content(table, start_date_row, start_date_col, start_date_content)
                    replace_specific_cell_content(table, venue_name_row, venue_name_col, venue_name_content)
                    replace_specific_cell_content(table, project_name_row, project_name_col, project_name_content)
                    replace_specific_cell_content(table, trigger_time_row, trigger_time_col, trigger_time_content)

                # 创建保存文件的文件夹
                folder_name = start_date_content.replace("/", "-")  # 将日期中的斜杠替换为横线，以便作为文件夹名称
                os.makedirs(os.path.join(output_folder, folder_name), exist_ok=True)

                # 保存修改后的Word文件到指定文件夹，文件名保持与输入的文件名一致
                output_file_path = os.path.join(output_folder, folder_name, file_name)
                doc.save(output_file_path)
                print(f"修改后的文件已保存为 {output_file_path}")
                break  # 找到文件后跳出循环
        else:
            print(f"未找到指定文件: {file_name}")


def main():
    # 用户输入多个文件名及其对应的发生时间，以横线分隔
    file_time_input = input("请输入多个文件名及其对应的发生时间（格式：文件名1-时间1,文件名2-时间2）：")
    file_time_pairs = [pair.strip() for pair in file_time_input.split(',')]

    # 构建文件名与发生时间的映射
    file_time_mapping = {}
    for pair in file_time_pairs:
        file_name, trigger_time_content = pair.split('-')
        file_time_mapping[file_name] = trigger_time_content

    # 用户输入开始日期的行列索引
    # start_date_row = int(input("请输入开始日期所在的行索引："))
    # start_date_col = int(input("请输入开始日期所在的列索引："))
    start_date_row = 1
    start_date_col = 1
    start_date_content = input("请输入需要修改的开始日期内容：")

    # 用户输入发生时间的行列索引
    # trigger_time_row = int(input("请输入发生时间所在的行索引："))
    # trigger_time_col = int(input("请输入发生时间所在的列索引："))
    trigger_time_row = 1
    trigger_time_col = 2

    # 用户输入场馆名称的行列索引
    # venue_name_row = int(input("请输入场馆名称所在的行索引："))
    # venue_name_col = int(input("请输入场馆名称所在的列索引："))
    venue_name_row = 1
    venue_name_col = 3
    venue_name_content = input("请输入需要修改的场馆名称内容：")

    # 用户输入项目名称的行列索引
    # project_name_row = int(input("请输入项目名称所在的行索引："))
    # project_name_col = int(input("请输入项目名称所在的列索引："))
    project_name_row = 1
    project_name_col = 4
    project_name_content = input("请输入需要修改的项目名称内容：")
    

    # 用户输入保存文件的文件夹名称
    output_folder = input("请输入保存文件的文件夹名称：")

    process_files(file_time_mapping, start_date_row, start_date_col, start_date_content,
                  trigger_time_row, trigger_time_col,
                  venue_name_row, venue_name_col, venue_name_content,
                  project_name_row, project_name_col, project_name_content, output_folder)


if __name__ == "__main__":
    main()
