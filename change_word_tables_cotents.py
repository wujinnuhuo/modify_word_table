import os
from docx import Document
from collections import defaultdict

def replace_specific_cell_content(table, row_index, col_index, new_content):
    if row_index < len(table.rows) and col_index < len(table.columns):
        cell = table.cell(row_index, col_index)
        cell.text = new_content


def main():
    root_folder = input("请输入文件夹路径: ")

    if not os.path.exists(root_folder):
        print("指定的文件夹路径不存在。")
        return

    file_dict = defaultdict(list)

    for folder, _, files in os.walk(root_folder):
        for filename in files:
            file_path = os.path.join(folder, filename)
            doc = Document(file_path)
            table = doc.tables[1]
            replace_specific_cell_content(table, 1, 2, "场景完成情况\nScenario situation and observation")
            replace_specific_cell_content(table, 1, 4, "其他：\nMiscellaneous")
            output_file_path = os.path.join(folder, filename)
            doc.save(output_file_path)

            
   

if __name__ == "__main__":
    main()